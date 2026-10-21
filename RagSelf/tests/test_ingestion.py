"""Run: python -m pytest -v tests/  (from directory containing ingestor.py)."""
import json
from pathlib import Path
from uuid import UUID

import pytest
from langchain_core.documents import Document

import ingestor as m
from conftest import write_txt


def assert_synced(engine):
    report = engine.validate_ingestion()
    assert report['missing_in_qdrant'] == [], report
    assert report['orphaned_in_qdrant'] == [], report
    assert report['missing_source_files'] == [], report
    assert report['modified_source_files'] == [], report
    assert report['manifest_points'] == report['qdrant_points'], report


def test_normalization_repairs_unicode_and_whitespace():
    result = m.normalize_text('  ﬁle  \r\n\r\n\r\n hello  world  ')
    assert result == 'file\n\nhello world'


def test_source_id_is_stable_and_path_dependent(tmp_path):
    root = tmp_path / 'data'
    a, b = root / 'a.txt', root / 'folder' / 'a.txt'
    a.parent.mkdir(); b.parent.mkdir()
    a.write_text('same', encoding='utf-8')
    b.write_text('same', encoding='utf-8')
    assert m.source_key(b, root) == 'folder/a.txt'
    assert m.source_id(a, root) == m.source_id(a, root)
    assert m.source_id(a, root) != m.source_id(b, root)
    assert m.file_hash(a) == m.file_hash(b)
    with pytest.raises(ValueError):
        m.source_key(tmp_path / 'outside.txt', root)


def test_txt_loader_metadata_and_chunk_ids(sandbox, sample_text):
    engine, data = sandbox
    path = write_txt(data, 'policy.txt', sample_text)
    docs = m.standardized_docs(path, data)
    assert len(docs) == 1
    assert docs[0].metadata['source_path'] == 'policy.txt'
    assert docs[0].metadata['file_type'] == '.txt'
    assert docs[0].metadata['page'] is None
    chunks = m.chunk_source(path, data, engine.splitter, True)
    assert chunks
    assert [c.id for c in chunks] == [c.id for c in m.chunk_source(path, data, engine.splitter, True)]
    assert len({c.id for c in chunks}) == len(chunks)
    assert [c.metadata['chunk_index'] for c in chunks] == list(range(len(chunks)))
    for c in chunks:
        assert UUID(m.point_id(c.id)).version == 5
        payload = m.make_payload(c)
        assert payload['text'] == c.page_content
        assert payload['content_hash'] == m.content_hash(c.page_content)
        assert payload['chunk_id'] == c.id


def test_pdf_loader_preserves_pages(sandbox):
    pytest.importorskip('reportlab')
    from reportlab.pdfgen import canvas
    engine, data = sandbox
    path = data / 'example.pdf'
    pdf = canvas.Canvas(str(path))
    pdf.drawString(50, 760, 'FIRST PAGE leave policy')
    pdf.showPage()
    pdf.drawString(50, 760, 'SECOND PAGE security policy')
    pdf.save()
    docs = m.standardized_docs(path, data)
    assert len(docs) == 2
    assert [d.metadata['page'] for d in docs] == [0, 1]
    assert 'FIRST PAGE' in docs[0].page_content
    assert 'SECOND PAGE' in docs[1].page_content


def test_docx_loader(sandbox):
    pytest.importorskip('docx')
    from docx import Document as WordDoc
    engine, data = sandbox
    path = data / 'terms.docx'
    doc = WordDoc(); doc.add_paragraph('Personnel policy document'); doc.save(str(path))
    docs = m.standardized_docs(path, data)
    assert 'Personnel policy' in docs[0].page_content
    assert docs[0].metadata['file_type'] == '.docx'


def test_unsupported_file_rejected(sandbox):
    engine, data = sandbox
    path = data / 'image.jpg'
    path.write_bytes(b'fake')
    with pytest.raises(ValueError, match='Unsupported'):
        engine.ingest_one(path)


def test_structure_aware_section_split(sandbox):
    engine, data = sandbox
    doc = Document(page_content='# First\nalpha paragraph\n\n## Second\nbeta paragraph', metadata={})
    sections = m._sections(doc, True)
    assert len(sections) >= 2
    assert any('First' in s.page_content for s in sections)
    assert any('Second' in s.page_content for s in sections)
    assert m._sections(doc, False) == [doc]


def test_new_ingestion_manifest_payload_and_validation(sandbox, sample_text):
    engine, data = sandbox
    path = write_txt(data, 'policy.txt', sample_text)
    result = engine.ingest_one(path)
    assert result['skipped'] is False
    assert result['chunks'] > 0
    manifest = m.existing_manifest(engine.conn)
    sid = m.source_id(path, data)
    assert manifest[sid]['hash'] == m.file_hash(path)
    assert len(manifest[sid]['ids']) == result['chunks']
    assert_synced(engine)
    points = list(engine.stored_chunks())
    assert all('chunk_id' in p.payload and 'content_hash' in p.payload for p in points)
    assert all(len(engine.client.retrieve(engine.config.COLLECTION, [p.id], with_vectors=True)[0].vector) == 8 for p in points)


def test_unchanged_document_skips_without_embedding(sandbox, sample_text):
    engine, data = sandbox
    path = write_txt(data, 'policy.txt', sample_text)
    engine.ingest_one(path)
    before = len(engine.embeddings.calls)
    assert engine.ingest_one(path) == {'skipped': True, 'reused': 0, 'embedded': 0}
    assert len(engine.embeddings.calls) == before
    assert_synced(engine)


def test_change_scan_detects_new_modified_and_removed(sandbox, sample_text):
    engine, data = sandbox
    p1 = write_txt(data, 'one.txt', sample_text)
    engine.scan_changes()
    assert p1 in engine.changed_files
    engine.ingest_one(p1)
    engine.scan_changes()
    assert not engine.changed_files and not engine.removed_ids
    p1.write_text(sample_text + '\nChange notice', encoding='utf-8')
    p2 = write_txt(data, 'two.txt', 'A new source')
    engine.scan_changes()
    assert set(engine.changed_files) == {p1, p2}
    p1.unlink()
    engine.scan_changes()
    assert m.source_id(p1, data) in engine.removed_ids


def test_modification_cleans_stale_points(sandbox, sample_text):
    engine, data = sandbox
    path = write_txt(data, 'changed.txt', sample_text)
    engine.ingest_one(path)
    previous = set(engine.manifest[m.source_id(path, data)]['ids'])
    path.write_text('Completely different short text.', encoding='utf-8')
    result = engine.ingest_one(path)
    current = set(engine.manifest[m.source_id(path, data)]['ids'])
    assert result['skipped'] is False
    assert current != previous
    existing = {str(p.id) for p in engine.stored_chunks()}
    assert all(m.point_id(cid) not in existing for cid in previous - current)
    assert_synced(engine)


def test_embedding_reuse_between_identical_sources(sandbox, sample_text):
    engine, data = sandbox
    first = write_txt(data, 'first.txt', sample_text)
    second = write_txt(data, 'second.txt', sample_text)
    engine.ingest_one(first)
    before = sum(len(call) for call in engine.embeddings.calls)
    result = engine.ingest_one(second)
    after = sum(len(call) for call in engine.embeddings.calls)
    assert result['reused'] > 0
    assert result['embedded'] == 0
    assert before == after
    assert_synced(engine)


def test_exact_duplicate_report_non_destructive(sandbox, sample_text):
    engine, data = sandbox
    p1 = write_txt(data, 'one.txt', sample_text)
    p2 = write_txt(data, 'two.txt', sample_text)
    engine.ingest_one(p1)
    candidate = m.chunk_source(p2, data, engine.splitter, engine.config.STRUCTURE_AWARE)
    before = len(list(engine.stored_chunks()))
    groups = engine.exact_duplicate_report(candidate)
    assert groups
    assert any({'stored', 'incoming'} <= {r[0] for r in refs} for refs in groups.values())
    assert len(list(engine.stored_chunks())) == before


def test_near_duplicates_exclude_exact(sandbox):
    engine, data = sandbox
    p = write_txt(data, 'a.txt', 'The company allows twenty annual leave days.')
    engine.ingest_one(p)
    almost = Document(page_content='The company allows twenty annual leave days!', id='candidate')
    exact = Document(page_content='The company allows twenty annual leave days.', id='same')
    assert engine.near_duplicate_report([almost], threshold=80)
    assert engine.near_duplicate_report([exact], threshold=80) == []


def test_remove_source_only_deletes_its_own_points(sandbox, sample_text):
    engine, data = sandbox
    a = write_txt(data, 'a.txt', sample_text)
    b = write_txt(data, 'b.txt', sample_text)
    engine.ingest_one(a); engine.ingest_one(b)
    sid_a, sid_b = m.source_id(a, data), m.source_id(b, data)
    engine.remove_source(sid_a)
    assert sid_a not in engine.manifest
    assert sid_b in engine.manifest
    # The original file still exists, so it is a new pending source rather than a missing file.
    assert_synced(engine)


def test_validation_detects_missing_source_modified_source_and_missing_point(sandbox, sample_text):
    engine, data = sandbox
    a = write_txt(data, 'a.txt', sample_text)
    b = write_txt(data, 'b.txt', 'other text')
    engine.ingest_one(a); engine.ingest_one(b)
    b.write_text('modified contents', encoding='utf-8')
    a.unlink()
    report = engine.validate_ingestion()
    assert 'a.txt' in report['missing_source_files']
    assert 'b.txt' in report['modified_source_files']
    cid = engine.manifest[m.source_id(b, data)]['ids'][0]
    m.delete_chunk_ids([cid], engine.client, engine.config.COLLECTION)
    report = engine.validate_ingestion()
    assert m.point_id(cid) in report['missing_in_qdrant']


def test_orphan_cleanup_dry_run_and_actual(sandbox):
    engine, data = sandbox
    from qdrant_client import models
    pid = '00000000-0000-4000-8000-000000000001'
    engine.client.upsert(
        collection_name=engine.config.COLLECTION,
        points=[models.PointStruct(id=pid, vector=[1.0] + [0.0]*7, payload={'text':'orphan'})],
        wait=True,
    )
    assert pid in engine.cleanup_orphans(dry_run=True)
    assert pid in {str(p.id) for p in engine.stored_chunks()}
    assert pid in engine.cleanup_orphans(dry_run=False)
    assert pid not in {str(p.id) for p in engine.stored_chunks()}



def test_interrupted_upsert_keeps_manifest_uncommitted(sandbox, sample_text, monkeypatch):
    engine, data = sandbox
    # Generate enough text to produce more than 64 chunks
    text = "\n\n".join(
        f"Section {i}: " + ("Unique policy content. " * 100)
        for i in range(100)
    )
    path = write_txt(data, "failure.txt", text)
    chunks = m.chunk_source(path, data, engine.splitter)
    assert len(chunks) > 64, "Test requires multiple upsert batches"
    original_upsert = engine.client.upsert
    calls = 0
    def fail_second_batch(*args, **kwargs):
        nonlocal calls
        calls += 1
        if calls == 2:
            raise RuntimeError("simulated interruption")
        return original_upsert(*args, **kwargs)
    # Simulate interruption during second Qdrant batch
    monkeypatch.setattr(
        engine.client, "upsert", fail_second_batch
    )

    with pytest.raises(RuntimeError, match="simulated interruption"):
        engine.ingest_one(path)
    # SQLite must not commit the failed ingestion
    sid = m.source_id(path, data)
    assert sid not in m.existing_manifest(engine.conn)
    # Staging should be cleared even after failure
    assert engine.staged == {}
    # Restore normal Qdrant behavior
    monkeypatch.setattr(
        engine.client, "upsert", original_upsert
    )
    # Retry ingestion
    result = engine.ingest_one(path)
    assert result["skipped"] is False
    # Verify SQLite and Qdrant are synchronized
    assert_synced(engine)


def test_manifest_chunk_ids_are_serialized_json(sandbox):
    engine, data = sandbox
    path = write_txt(data, 'json.txt', 'Manifest check')
    engine.ingest_one(path)
    (raw,) = engine.conn.execute('SELECT chunk_ids_json FROM sources').fetchone()
    assert isinstance(raw, str)
    assert isinstance(json.loads(raw), list)
    assert json.loads(raw) == engine.manifest[m.source_id(path, data)]['ids']
